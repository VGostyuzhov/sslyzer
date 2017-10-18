from openpyxl import Workbook
from openpyxl.styles import NamedStyle
from styles import styles

def excel_report(servers, filename):
    try:
        wb = Workbook()
    except Exception, error:
        raise error
    for key, value in styles.iteritems():
        wb.add_named_style(value)

    sheet = wb.create_sheet('SSL')
    headers = {'Domain': 1, 'SSLv2': 2, 'SSLv3': 3, 'TLSv1.0': 4, 'TLSv1.1': 5, 'TLSv1.2': 6,
            'Heartbleed': 7, 'Crime': 8, 'Downgrade': 9, 'Poodle': 10 , 'RC4': 11, 'Beast': 12, 'CCS Injection': 13, 'Drown': 14, 'Freak': 15, 'Logjam': 16,
            'Trusted': 17, 'Self Signed': 18, 'Valid to': 19, 'Matches hostname': 20}
    for key, value in headers.iteritems():
        sheet.cell(column=value, row=1, value=key)
        sheet.cell(column=value, row=1).style = styles['Header']
    row = 2
    for server in servers:
        if 'hostname' in server:
            sheet.cell(column=1, row=row, value=server['hostname'])
            sheet.cell(column=2, row=row, value=str(server['sslv20']))
            sheet.cell(column=3, row=row, value=str(server['sslv30']))
            sheet.cell(column=4, row=row, value=str(server['tlsv10']))
            sheet.cell(column=5, row=row, value=str(server['tlsv11']))
            sheet.cell(column=6, row=row, value=str(server['tlsv12']))
            sheet.cell(column=7, row=row, value=str(server['vulners']['heartbleed']))
            sheet.cell(column=8, row=row, value=str(server['vulners']['crime']))
            sheet.cell(column=9, row=row, value=str(server['vulners']['downgrade']))
            sheet.cell(column=10, row=row, value=str(server['vulners']['poodle']))
            sheet.cell(column=11, row=row, value=str(server['vulners']['rc4']))
            sheet.cell(column=12, row=row, value=str(server['vulners']['beast']))
            sheet.cell(column=13, row=row, value=str(server['vulners']['ccs_injection']))
            sheet.cell(column=14, row=row, value=str(server['vulners']['drown']))
            sheet.cell(column=15, row=row, value=str(server['vulners']['freak']))
            sheet.cell(column=16, row=row, value=str(server['vulners']['logjam']))
            sheet.cell(column=17, row=row, value=str(server['cert']['trusted']))
            sheet.cell(column=18, row=row, value=str(server['cert']['self_signed']))
            sheet.cell(column=19, row=row, value=str(server['cert']['not_valid_after']))
            sheet.cell(column=20, row=row, value=str(server['cert']['matches_hostname']))
            row += 1

    width_dict = {'A': 15, 'B': 7, 'C': 7, 'D': 7, 'E': 7, 'F': 7, 'G': 10}
    for key, value in width_dict.iteritems():
        sheet.column_dimensions[key].width = value

    try:
        wb.remove_sheet(wb.get_sheet_by_name('Sheet'))
        wb.save(filename)
    except Exception, error:
        raise error